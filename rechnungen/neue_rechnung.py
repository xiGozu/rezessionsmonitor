#!/usr/bin/env python3
"""Erstellt eine neue Rechnung aus vorlage.docx und zählt die Rechnungsnummer automatisch hoch.

Beispiel:
  python rechnungen/neue_rechnung.py \
    --empfaenger "Fa. Exakt|Ingolf Liebsch|Am Jacobstein 1|01445 Radebeul" \
    --leistungsort "Wohnhaus|Burgstädt" \
    --pos "103,5;m²;Parkettboden, Sockel (21.09.-02.10.2026);36"
(--pos: Anzahl;Einheit;Bezeichnung;Einzelpreis, mehrfach möglich)
"""
import argparse, copy, datetime
from pathlib import Path
import docx

HERE = Path(__file__).parent
MWST = 0.19


def num(s):
    return float(s.replace(".", "").replace(",", "."))


def eur(x):
    return f"{x:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".") + " €"


def fmt_qty(x):
    return f"{x:g}".replace(".", ",")


def set_text(par, text):
    """Text in den ersten Run schreiben (Formatierung bleibt), übrige Runs leeren; \\n = Zeilenumbruch."""
    runs = par.runs
    for r in runs[1:]:
        r._r.getparent().remove(r._r)
    r = runs[0]
    for t in list(r._r):
        if t.tag.endswith("}t") or t.tag.endswith("}br"):
            r._r.remove(t)
    for i, line in enumerate(text.split("\n")):
        if i:
            r.add_break()
        r.add_text(line)


def main():
    a = argparse.ArgumentParser()
    a.add_argument("--empfaenger", required=True, help="Zeilen mit | getrennt")
    a.add_argument("--leistungsort", required=True, help="Zeilen mit | getrennt")
    a.add_argument("--pos", action="append", required=True)
    a.add_argument("--datum", default=datetime.date.today().strftime("%d.%m.%Y"))
    a.add_argument("--zahlungsziel", default="7 Tage")
    a = a.parse_args()

    counter = HERE / "letzte_nummer.txt"
    nr = str(int(counter.read_text().strip()) + 1)

    d = docx.Document(HERE / "vorlage.docx")
    head, items = d.tables
    set_text(head.rows[1].cells[0].paragraphs[0], a.empfaenger.replace("|", "\n"))
    set_text(head.rows[1].cells[1].paragraphs[0], f"Datum: {a.datum}\nRechnung-Nr.: {nr}")
    set_text(d.paragraphs[3], a.leistungsort.replace("|", "\n"))
    zb = d.paragraphs[5].runs[-1]
    zb.text = f": {a.zahlungsziel}"

    # Positionszeilen: Zeile 1 als Muster, überzählige entfernen
    tmpl = copy.deepcopy(items.rows[1]._tr)
    items._tbl.remove(items.rows[1]._tr)
    items._tbl.remove(items.rows[1]._tr)
    anchor = items.rows[0]._tr
    netto = 0.0
    for i, p in enumerate(a.pos, 1):
        q, unit, text, price = p.split(";")
        q, price = num(q), num(price)
        tot = round(q * price, 2)
        netto += tot
        tr = copy.deepcopy(tmpl)
        anchor.addnext(tr)
        anchor = tr
        row = [r for r in items.rows if r._tr is tr][0]
        for c, v in zip(row.cells, [str(i), fmt_qty(q), unit, text, eur(price), eur(tot)]):
            set_text(c.paragraphs[0], v)
    netto = round(netto, 2)
    mwst = round(netto * MWST, 2)
    for row, v in zip(items.rows[-3:], [netto, mwst, round(netto + mwst, 2)]):
        set_text(row.cells[5].paragraphs[0], eur(v))

    out = HERE / "ausgang" / f"Rechnung_{nr}.docx"
    d.save(out)
    counter.write_text(nr + "\n")
    print(f"{out}  (netto {eur(netto)}, MwSt {eur(mwst)}, brutto {eur(netto + mwst)})")


if __name__ == "__main__":
    main()
